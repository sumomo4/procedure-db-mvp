"""UNI naming rules, independent target locations and immutable generation context."""

from io import BytesIO
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from app.core.case_doc_generation import _build_placeholder_values_for_group, build_case_doc_workbook_bytes
from app.core.responses import CaseDocDataSourceListData, CaseDocResolveContextRequest, CaseDocUniWorkSelection
from app.db import case_doc_data_sources as sources
from app.db.case_doc_uni import UniCaseDocMasterRepository
from app.db.case_doc_workflows import UniBuilding, generate_uni_hosts, load_uni_master, read_workflow
from app.db.case_docs import resolve_case_doc_context
from app.routers.health import get_app_settings
from test_case_docs import _fake_source_doc_detail


def book(path, rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


@pytest.fixture
def uni_settings(tmp_path, test_settings, monkeypatch):
    exports = tmp_path / "exports" / "uni"
    exports.mkdir(parents=True)
    config = tmp_path / "config"
    profiles = config / "case_doc_workflows"
    profiles.mkdir(parents=True)
    mapping = config / "placeholder_mapping.yml"
    mapping.write_text("version: 1\nplaceholders: []\n", encoding="utf-8")
    specific = config / "placeholder_mappings" / "uni.yml"
    specific.parent.mkdir()
    specific.write_text(yaml.safe_dump({"version": 1, "placeholders": [{"name": "UNI_SSH_IP", "enabled": True,
        "scope": "device", "device_type": "UNI", "source_file": "装置追加.xlsx", "key_column": "ホスト名", "value_column": "SSH", "source_column": "ssh"}]}), encoding="utf-8")
    (profiles / "uni.yml").write_text("version: 1\nmode: uni\ndefinition_file: uni.xlsx\n", encoding="utf-8")
    master = Workbook()
    sheet = master.active
    sheet.title = "定義"
    for col, value in enumerate(["ブロック", "府県", "都道府県コード", "設置ビル", "ホスト名規則", "ペアビル", "ホスト名規則"], 2):
        sheet.cell(3, col, value)
    locations = [UniBuilding("関西1", "京都", "京都テスト", "test26", "", ""), UniBuilding("関西2", "兵庫", "神戸テスト", "test28", "", "")]
    for index, location in enumerate(locations, 4):
        for col, value in enumerate([location.block, location.prefecture, 1, location.building, location.prefix, "", ""], 2):
            sheet.cell(index, col, value)
        sheet.cell(index, 19, location.block)
        for number in range(1, 35):
            sheet.cell(3, 19 + number, number)
            sheet.cell(index, 19 + number, location.prefecture)
    sheet.cell(3, 19, "収容エリア")
    master.save(profiles / "uni.xlsx")
    master.close()
    definition_rows = [["設置エリア", "設置府県", "設置ビル", "ホスト名", "TTSホスト名①", "TTSホスト名②"]]
    device_rows = [["ホスト名", "SSH"]]
    tts_rows = [["TTSホスト名", "TTSアドレス"]]
    for index, location in enumerate(locations, 1):
        for device_index, device in enumerate(generate_uni_hosts(location, 4), 10):
            definition_rows.append([location.block, location.prefecture, location.building, device.host_name, f"tts-test-{index}-0", f"tts-test-{index}-1"])
            device_rows.append([device.host_name, f"192.0.{index}.{device_index}"])
        tts_rows.extend([[f"tts-test-{index}-0", f"192.0.{index}.200"], [f"tts-test-{index}-1", f"192.0.{index}.201"]])
    book(exports / "TTS_定義.xlsx", definition_rows)
    book(exports / "装置追加.xlsx", device_rows)
    book(exports / "TTS.xlsx", tts_rows)
    book(exports / "収容エリア.xlsx", [["収容エリア", "1CL"], ["関西1", "京都"]])
    settings = test_settings.model_copy(update={"case_doc_master_source": "export_file", "case_doc_access_export_dir": str(exports.parent), "case_doc_placeholder_mapping_path": str(mapping)})
    source = sources._to_data(settings, {"source_id": "uni", "name": "UNI", "enabled": True, "description": ""})
    monkeypatch.setattr(sources, "list_data_sources", lambda *_args, **_kwargs: CaseDocDataSourceListData(items=[source]))
    return settings


def repository(settings):
    scoped, _ = sources.select_data_source(settings, "uni")
    return UniCaseDocMasterRepository(scoped.case_doc_access_export_dir, scoped.case_doc_placeholder_mapping_path, scoped.case_doc_workflow_path)


@pytest.mark.parametrize("block,cluster,usbe,usbh,control", [
    ("関西2", 4, "100000", "140000", "usce000200"),
    ("関西2", 16, "400000", "440000", "usce000800"),
    ("関西2", 17, "500000", "510000", "usle000900"),
    ("関西2", 18, "500000", "520000", "usle000900"),
    ("関西2", 19, "500000", "530000", "usle000900"),
    ("関西2", 20, "500000", "540000", "usle001000"),
    ("関西2", 21, "500000", "550000", "usle001000"),
    ("関西2", 23, "600000", "610000", "usle001100"),
    ("関西1", 19, "500000", "530000", "usce001000"),
    ("関西1", 21, "600000", "610000", "usce001100"),
])
def test_excel_hostname_rules_at_boundaries(block, cluster, usbe, usbh, control):
    result = generate_uni_hosts(UniBuilding(block, "test", "test", "prefix", "", ""), cluster)
    assert result[0].host_name == "prefixusbe" + usbe
    assert result[2].host_name == "prefixusbh" + usbh
    assert result[-1].host_name == "prefix" + control


@pytest.mark.parametrize("cluster", [0, 37, -1, 1.5, True])
def test_invalid_cluster_cannot_generate_a_hostname(cluster):
    with pytest.raises(ValueError):
        generate_uni_hosts(UniBuilding("関西2", "test", "test", "prefix", "", ""), cluster)


def test_uni_profile_does_not_need_a_unit_file(uni_settings):
    source = sources.list_data_sources(uni_settings).items[0]
    assert source.generation_mode == "uni" and source.generation_ready and source.file_count == 4
    repo = repository(uni_settings)
    assert [item.value for item in repo.list_blocks().items] == ["関西1", "関西2"]
    assert not repo.list_unit_configs(None, "京都テスト").items
    assert not (repo.export_dir / "ユニット構成.xlsx").exists()
    assert len(repo.list_clusters("関西2", "神戸テスト").items) == 34
    assert len(repo.list_placeholder_source_files().items) == 4


def selection(block="関西1", building="京都テスト"):
    return {"block": block, "building": building, "cluster_no": 4, "device_keys": ["USBE_0"]}


def test_different_buildings_are_resolved_and_snapshotted_per_target(uni_settings):
    payload = CaseDocResolveContextRequest(data_source_id="uni", source_doc_id=1,
        uni_targets=[selection(), selection("関西2", "神戸テスト"), selection()])
    context = resolve_case_doc_context(uni_settings, payload)
    assert context.unit_config is None and context.generation_mode == "uni"
    assert [item.building for item in context.target_contexts] == ["京都テスト", "神戸テスト", "京都テスト"]
    assert [group[0].slot_key for group in context.target_groups] == ["USBE_T1_0", "USBE_T2_0", "USBE_T3_0"]
    assert context.target_groups[0][0].host_name == context.target_groups[2][0].host_name
    assert context.work_summary["prefecture"] == "京都 / 兵庫"
    values = [_build_placeholder_values_for_group(context, group)[0] for group in context.target_groups]
    assert [value["UNI_SSH_IP"] for value in values] == ["192.0.1.10", "192.0.2.10", "192.0.1.10"]
    assert [value["TARGET_BUILDING"] for value in values] == ["京都テスト", "神戸テスト", "京都テスト"]
    assert len(context.data_source_revision) == 64


def test_multiple_devices_and_tts_candidates_keep_existing_group_semantics(uni_settings):
    payload = CaseDocResolveContextRequest(data_source_id="uni", source_doc_id=1,
        uni_targets=[{**selection(), "device_keys": ["USBE_0", "USBH_0"]}])
    context = resolve_case_doc_context(uni_settings, payload)
    values, ambiguous = _build_placeholder_values_for_group(context, context.target_groups[0])
    assert "UNI_SSH_IP" in ambiguous and "UNI_SSH_IP" not in values
    assert values["TARGET_BUILDING"] == "京都テスト"
    assert len(values["TARGET_DEVICE_HOSTNAME"].splitlines()) == 2
    candidates = repository(uni_settings).target_candidates(CaseDocUniWorkSelection(**{key: value for key, value in selection().items() if key != "device_keys"}))
    assert len([item for item in candidates.devices if item.device_type == "TTS" and item.available]) == 2


def test_generated_excel_includes_target_location_snapshot(uni_settings):
    context = resolve_case_doc_context(uni_settings, CaseDocResolveContextRequest(data_source_id="uni", source_doc_id=1,
        uni_targets=[selection(), selection("関西2", "神戸テスト")]))
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{UNI_SSH_IP}} {{TARGET_BUILDING}}"
    data = build_case_doc_workbook_bytes(context, source)
    workbook = load_workbook(BytesIO(data), keep_vba=True)
    strings = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
    assert "targetごとの対象工事" in strings
    assert any("192.0.1.10 京都テスト" in value for value in strings)
    assert any("192.0.2.10 神戸テスト" in value for value in strings)
    workbook.close()


@pytest.mark.parametrize("change", ["missing_host", "duplicate_host", "wrong_building"])
def test_bad_export_data_is_not_selectable(uni_settings, change):
    repo = repository(uni_settings)
    path = repo.export_dir / "TTS_定義.xlsx"
    workbook = load_workbook(path)
    if change == "missing_host":
        workbook.active.delete_rows(2)
    elif change == "duplicate_host":
        workbook.active.append([cell.value for cell in workbook.active[2]])
    else:
        workbook.active.cell(2, 3, "別のビル")
    workbook.save(path)
    workbook.close()
    candidate = repo.target_candidates(CaseDocUniWorkSelection(block="関西1", building="京都テスト", cluster_no=4))
    assert not next(item for item in candidate.devices if item.slot_key == "USBE_0").available
    with pytest.raises(ValueError, match="照合できません"):
        repo.resolve_context(CaseDocResolveContextRequest(source_doc_id=1, uni_targets=[selection()]))


@pytest.mark.parametrize("bad", [{"building": "other"}, {"cluster_no": 35}, {"device_keys": ["USBE_0", "USBE_0"]}, {"device_keys": ["arbitrary-host"]}])
def test_forged_or_unavailable_selections_are_rejected(uni_settings, bad):
    with pytest.raises(ValueError):
        repository(uni_settings).resolve_context(CaseDocResolveContextRequest(source_doc_id=1, uni_targets=[{**selection(), **bad}]))


def test_rule_changes_affect_revision_but_not_the_resolved_snapshot(uni_settings):
    context = resolve_case_doc_context(uni_settings, CaseDocResolveContextRequest(data_source_id="uni", source_doc_id=1, uni_targets=[selection()]))
    frozen = context.model_dump_json()
    scoped, _ = sources.select_data_source(uni_settings, "uni")
    path = Path(scoped.case_doc_workflow_path).with_suffix(".xlsx")
    workbook = load_workbook(path)
    workbook["定義"].cell(4, 6, "changed")
    workbook.save(path)
    workbook.close()
    assert sources.data_source_revision(scoped) != context.data_source_revision
    assert context.model_dump_json() == frozen
    with pytest.raises(ValueError, match="再確認"):
        resolve_case_doc_context(uni_settings, CaseDocResolveContextRequest(data_source_id="uni", source_doc_id=1,
            uni_targets=[selection()], expected_data_source_revision=context.data_source_revision))


@pytest.mark.parametrize("name", ["../elsewhere.xlsx", "C:\\elsewhere.xlsx", "/tmp/elsewhere.xlsx"])
def test_profile_cannot_reference_a_master_outside_its_directory(tmp_path, name):
    profile = tmp_path / "uni.yml"
    profile.write_text(yaml.safe_dump({"version": 1, "mode": "uni", "definition_file": name}), encoding="utf-8")
    with pytest.raises(ValueError):
        read_workflow(profile)


def test_uni_api_candidates_and_preview_context_are_source_scoped(client, uni_settings, monkeypatch):
    client.app.dependency_overrides[get_app_settings] = lambda: uni_settings
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda settings, source_doc_id: _fake_source_doc_detail(source_doc_id))
    response = client.post("/api/v1/case-docs/uni/target-candidates?data_source_id=uni", json={"block": "関西1", "building": "京都テスト", "cluster_no": 4})
    assert response.status_code == 200
    assert response.json()["data"]["work"]["hostname_prefix"] == "test26"
    response = client.post("/api/v1/case-docs/resolve-context", json={"data_source_id": "uni", "source_doc_id": 1, "uni_targets": [selection()]})
    assert response.status_code == 200 and response.json()["data"]["unit_config"] is None
    response = client.get("/api/v1/case-docs/uni/clusters?data_source_id=uni&block=関西1&building=京都テスト")
    assert response.status_code == 200 and len(response.json()["data"]["items"]) == 34
    payload = {"data_source_id": "uni", "source_doc_id": 1, "uni_targets": [selection(), selection("関西2", "神戸テスト")]}
    response = client.post("/api/v1/case-docs/preview", json=payload)
    assert response.status_code == 200 and response.json()["data"]["context"]["generation_mode"] == "uni"
    response = client.post("/api/v1/case-docs/generate", json=payload)
    assert response.status_code == 200 and 'filename="case-doc-1-uni.xlsm"' in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), keep_vba=True)
    assert workbook.vba_archive is not None
    workbook.close()
