"""Device placeholders resolve with each definition's own lookup key."""

from io import BytesIO
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from app.core.case_doc_generation import build_case_doc_workbook_bytes
from app.core.responses import CaseDocResolveContextRequest, ModuleRowDeviceEntryData
from app.db.case_doc_repositories import ExportFileCaseDocMasterRepository
from test_case_docs import _fake_source_doc_detail


UNIT_FILE = "\u30e6\u30cb\u30c3\u30c8\u69cb\u6210.xlsx"
GUI_0_COLUMN = "GUI_0\u7cfb"
GUI_1_COLUMN = "GUI_1\u7cfb"
HOST_COLUMN = "\u30db\u30b9\u30c8\u540d"


def _write_workbook(path: Path, rows: list[list[str]]) -> None:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


def _mapping(name: str, key_column: str, value_column: str, **overrides) -> dict:
    return {
        "name": name,
        "enabled": True,
        "scope": "device",
        "device_type": "GUI",
        "source_file": UNIT_FILE,
        "key_column": key_column,
        "value_column": value_column,
        "source_column": name.lower(),
        **overrides,
    }


def _repository(tmp_path: Path, mappings: list[dict]) -> ExportFileCaseDocMasterRepository:
    _write_workbook(tmp_path / UNIT_FILE, [
        ["fs_cluster_name", "block", "prefecture", "building", GUI_0_COLUMN, GUI_1_COLUMN],
        ["FS-cluster", "block-1", "Tokyo", "test-building", "GUI-A", "GUI-B"],
    ])
    mapping_path = tmp_path / "placeholder_mapping.yml"
    mapping_path.write_text(yaml.safe_dump({"placeholders": mappings}), encoding="utf-8")
    return ExportFileCaseDocMasterRepository(str(tmp_path), str(mapping_path))


def _context(repository: ExportFileCaseDocMasterRepository, groups=None):
    return repository.resolve_context(CaseDocResolveContextRequest(
        source_doc_id=1,
        prefecture="Tokyo",
        building="test-building",
        block="block-1",
        target_slot_groups=groups or [["GUI_0"], ["GUI_1"]],
    ))


@pytest.fixture(params=[False, True], ids=["gui0-first", "gui1-first"])
def gui_repository(tmp_path: Path, request) -> ExportFileCaseDocMasterRepository:
    mappings = [
        _mapping("UNIT_GUI_0_HOSTNAME", GUI_0_COLUMN, GUI_0_COLUMN),
        _mapping("UNIT_GUI_1_HOSTNAME", GUI_1_COLUMN, GUI_1_COLUMN),
    ]
    if request.param:
        mappings.reverse()
    return _repository(tmp_path, mappings)


def test_gui_definitions_resolve_independent_host_names(gui_repository) -> None:
    context = _context(gui_repository)
    assert {
        (item.placeholder, item.host_name): item.value
        for item in context.resolved_placeholders
    } == {
        ("UNIT_GUI_0_HOSTNAME", "GUI-A"): "GUI-A",
        ("UNIT_GUI_1_HOSTNAME", "GUI-B"): "GUI-B",
    }


def test_gui_definitions_replace_both_target_columns_in_excel(gui_repository) -> None:
    context = _context(gui_repository)
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].device_entries = [
        ModuleRowDeviceEntryData(slot_no=1, command_text="{{UNIT_GUI_0_HOSTNAME}}"),
        ModuleRowDeviceEntryData(slot_no=2, command_text="{{UNIT_GUI_1_HOSTNAME}}"),
    ]
    source.items[0].rows[1].command_text = "{{TARGET_DEVICE_HOSTNAME}}"
    workbook = load_workbook(BytesIO(build_case_doc_workbook_bytes(context, source)), keep_vba=True)
    try:
        sheet = workbook.worksheets[0]
        assert [sheet[cell].value for cell in ("M4", "Q4")] == ["GUI-A", "GUI-B"]
        assert [sheet[cell].value for cell in ("M6", "Q6")] == ["GUI-A", "GUI-B"]
        assert [sheet[cell].value for cell in ("M7", "Q7")] == ["GUI-A", "GUI-B"]
    finally:
        workbook.close()


def test_gui_definitions_replace_both_members_in_one_target(gui_repository) -> None:
    context = _context(gui_repository, [["GUI_0", "GUI_1"]])
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{UNIT_GUI_0_HOSTNAME}} / {{UNIT_GUI_1_HOSTNAME}}"
    source.items[0].rows[1].command_text = "{{TARGET_DEVICE_HOSTNAME}}"
    workbook = load_workbook(BytesIO(build_case_doc_workbook_bytes(context, source)), keep_vba=True)
    try:
        sheet = workbook.worksheets[0]
        assert sheet["M6"].value == "GUI-A / GUI-B"
        assert sheet["M7"].value == "GUI-A\nGUI-B"
    finally:
        workbook.close()


def test_gui_definition_is_not_replaced_in_the_other_devices_target(gui_repository) -> None:
    context = _context(gui_repository)
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].device_entries = [
        ModuleRowDeviceEntryData(slot_no=1, command_text="{{UNIT_GUI_1_HOSTNAME}}"),
        ModuleRowDeviceEntryData(slot_no=2, command_text="{{UNIT_GUI_0_HOSTNAME}}"),
    ]
    workbook = load_workbook(BytesIO(build_case_doc_workbook_bytes(context, source)), keep_vba=True)
    try:
        sheet = workbook.worksheets[0]
        assert sheet["M6"].value == "{{UNIT_GUI_1_HOSTNAME}}"
        assert sheet["Q6"].value == "{{UNIT_GUI_0_HOSTNAME}}"
    finally:
        workbook.close()


def test_explicit_key_is_not_overridden_by_a_generic_host_column(tmp_path: Path) -> None:
    repository = _repository(tmp_path, [
        _mapping("GUI_VALUE", "GUI lookup", "value", source_file="GUI.xlsx"),
    ])
    _write_workbook(tmp_path / "GUI.xlsx", [
        ["host_name", HOST_COLUMN, " GUI lookup ", "value"],
        ["unrelated-host", "another-host", "GUI-B", "correct-value"],
    ])
    assert [(item.host_name, item.value) for item in _context(repository).resolved_placeholders] == [
        ("GUI-B", "correct-value"),
    ]


@pytest.mark.parametrize("empty_key", ["missing", "blank"])
def test_unavailable_key_does_not_use_another_definitions_key(tmp_path: Path, empty_key: str) -> None:
    repository = _repository(tmp_path, [
        _mapping("GUI_FIRST", "first_key", "first_value", source_file="GUI.xlsx"),
        _mapping("GUI_SECOND", "second_key", "second_value", source_file="GUI.xlsx"),
    ])
    columns = ["host_name", "second_key", "first_value", "second_value"]
    values = ["GUI-A", "GUI-B", "must-not-resolve", "second-value"]
    if empty_key == "blank":
        columns.append("first_key")
        values.append("")
    _write_workbook(tmp_path / "GUI.xlsx", [columns, values])
    assert [(item.placeholder, item.host_name, item.value) for item in _context(repository).resolved_placeholders] == [
        ("GUI_SECOND", "GUI-B", "second-value"),
    ]


@pytest.mark.parametrize("separate_files", [False, True], ids=["same-file", "separate-files"])
def test_definitions_do_not_share_values_by_internal_key(tmp_path: Path, separate_files: bool) -> None:
    second_file = "GUI_extra.xlsx" if separate_files else "GUI.xlsx"
    repository = _repository(tmp_path, [
        _mapping("GUI_FIRST", "host_name", "first_value", source_file="GUI.xlsx", source_column="shared_key"),
        _mapping("GUI_SECOND", "host_name", "second_value", source_file=second_file, source_column="shared_key"),
    ])
    _write_workbook(tmp_path / "GUI.xlsx", [
        ["host_name", "first_value", "second_value"], ["GUI-A", "first-value", "second-value"],
    ])
    if separate_files:
        _write_workbook(tmp_path / second_file, [["host_name", "second_value"], ["GUI-A", "second-value"]])
    assert {item.placeholder: item.value for item in _context(repository).resolved_placeholders} == {
        "GUI_FIRST": "first-value", "GUI_SECOND": "second-value",
    }


@pytest.mark.parametrize("key_column, header", [(HOST_COLUMN, "host_name"), ("host_name", HOST_COLUMN)])
def test_host_name_header_aliases_remain_supported(tmp_path: Path, key_column: str, header: str) -> None:
    repository = _repository(tmp_path, [_mapping("GUI_VALUE", key_column, "value", source_file="GUI.xlsx")])
    _write_workbook(tmp_path / "GUI.xlsx", [[header, "value"], ["GUI-A", "resolved-value"]])
    assert [(item.host_name, item.value) for item in _context(repository).resolved_placeholders] == [
        ("GUI-A", "resolved-value"),
    ]


@pytest.mark.parametrize("key_column, alias", [(HOST_COLUMN, "host_name"), ("host_name", HOST_COLUMN)])
def test_configured_host_header_takes_precedence_over_its_alias(tmp_path: Path, key_column: str, alias: str) -> None:
    repository = _repository(tmp_path, [_mapping("GUI_VALUE", key_column, "value", source_file="GUI.xlsx")])
    _write_workbook(tmp_path / "GUI.xlsx", [[key_column, alias, "value"], ["GUI-B", "GUI-A", "resolved-value"]])
    assert [(item.host_name, item.value) for item in _context(repository).resolved_placeholders] == [
        ("GUI-B", "resolved-value"),
    ]


def test_present_but_blank_host_key_does_not_fall_back_to_an_alias(tmp_path: Path) -> None:
    repository = _repository(tmp_path, [_mapping("GUI_VALUE", HOST_COLUMN, "value", source_file="GUI.xlsx")])
    _write_workbook(tmp_path / "GUI.xlsx", [[HOST_COLUMN, "host_name", "value"], ["", "GUI-A", "wrong-value"]])
    assert _context(repository).resolved_placeholders == []
