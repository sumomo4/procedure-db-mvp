"""Execution indentation is snapshotted and legacy cases use their own workbook."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
import pytest
from pydantic import ValidationError

from app.core.case_doc_generation import (
    _read_case_doc_execution_indents,
    build_case_doc_execution_item_snapshots,
    build_case_doc_workbook_bytes,
    read_case_doc_execution_indents,
)
from app.core.responses import CaseDocExecutionItemData
from test_case_docs import _fake_case_doc_context, _fake_source_doc_detail


def make_workbook(level: int | None, path: Path) -> tuple[bytes, list[dict]]:
    source = _fake_source_doc_detail(1)
    for row in source.items[0].rows:
        row.indent_level = level
    context = _fake_case_doc_context()
    content = build_case_doc_workbook_bytes(context, source)
    path.write_bytes(content)
    return content, build_case_doc_execution_item_snapshots(source, context, content)


@pytest.mark.parametrize("level", [None, 0, 1, 2, 3])
def test_snapshot_and_legacy_original_recover_the_same_indent(tmp_path, level):
    path = tmp_path / "original.xlsm"
    content, snapshots = make_workbook(level, path)
    assert len(snapshots) == 4
    assert [item["indent_level"] for item in snapshots] == [level or 0] * 4
    original_hash = sha256(content).hexdigest()
    levels = read_case_doc_execution_indents(path)
    workbook = load_workbook(BytesIO(content), keep_vba=True)
    try:
        for item in snapshots:
            row = workbook.worksheets[0][item["excel_cell"]].row
            assert levels[row] == item["indent_level"]
    finally:
        workbook.close()
    assert sha256(path.read_bytes()).hexdigest() == original_hash


def test_snapshot_remains_unchanged_after_source_module_edit(tmp_path):
    source = _fake_source_doc_detail(1)
    context = _fake_case_doc_context()
    source.items[0].rows[0].indent_level = 2
    content = build_case_doc_workbook_bytes(context, source)
    snapshots = build_case_doc_execution_item_snapshots(source, context, content)
    path = tmp_path / "original.xlsm"
    path.write_bytes(content)
    source.items[0].rows[0].indent_level = 0
    assert snapshots[0]["indent_level"] == 2
    levels = read_case_doc_execution_indents(path)
    assert levels[int(snapshots[0]["excel_cell"][1:])] == 2


def test_legacy_indent_supports_column_position_and_excel_cell_alignment(tmp_path):
    path = tmp_path / "original.xlsm"
    content, snapshots = make_workbook(1, path)
    workbook = load_workbook(BytesIO(content), keep_vba=True)
    try:
        row = int(snapshots[0]["excel_cell"][1:])
        workbook.worksheets[0].cell(row=row, column=6).alignment = Alignment(indent=1)
        workbook.save(path)
    finally:
        workbook.close()
    assert read_case_doc_execution_indents(path)[row] == 2


def test_legacy_indent_is_cached_and_invalidated_when_original_changes(tmp_path, monkeypatch):
    path = tmp_path / "original.xlsm"
    _, snapshots = make_workbook(1, path)
    _read_case_doc_execution_indents.cache_clear()
    first = read_case_doc_execution_indents(path)
    assert _read_case_doc_execution_indents.cache_info().misses == 1
    row = int(snapshots[0]["excel_cell"][1:])
    first[row] = 3
    assert read_case_doc_execution_indents(path)[row] == 1
    assert _read_case_doc_execution_indents.cache_info().hits == 1
    make_workbook(3, path)
    assert read_case_doc_execution_indents(path)[row] == 3
    assert _read_case_doc_execution_indents.cache_info().misses == 2
    _read_case_doc_execution_indents.cache_clear()


def test_missing_original_does_not_fall_back_to_current_module(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_case_doc_execution_indents(tmp_path / "missing.xlsm")


@pytest.mark.parametrize("value", [-1, 4])
def test_execution_indent_is_limited_to_supported_excel_columns(value):
    with pytest.raises(ValidationError):
        CaseDocExecutionItemData(execution_item_id=1, row_order=1, target_no=1,
            host_name="gui-test-0", excel_cell="J6", status="pending", lock_version=0, indent_level=value)


def test_older_api_payloads_default_to_zero_indent():
    item = CaseDocExecutionItemData(execution_item_id=1, row_order=1, target_no=1,
        host_name="gui-test-0", excel_cell="J6", status="pending", lock_version=0)
    assert item.indent_level == 0
