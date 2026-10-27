"""Generation previews share Excel substitution rules and never persist cases."""

from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.case_doc_generation import TEMPLATE_PATH
from app.core.security import get_current_user
from app.core.exceptions import DatabaseConnectionError
from test_case_doc_target_groups import _export_context, _request
from test_case_docs import _fake_source_doc_detail, _mock_source_doc_detail


def test_preview_matches_download_and_does_not_create_an_instance(client, monkeypatch):
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].work_text = "Check {{UNKNOWN_DEVICE}} / {{UNKNOWN_DEVICE}}"
    source.items[0].rows[1].command_text = "{{TARGET_DEVICE_HOSTNAME}}"
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda *_: source)
    monkeypatch.setattr("app.routers.case_docs.create_case_doc_instance", lambda *_: pytest.fail("Preview must not persist"))
    before = sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
    payload = _request(client, [["SBC_CL1_0"], ["SBC_CL1_1"], ["SBC_CL1_0"]])
    preview_response = client.post("/api/v1/case-docs/preview", json=payload)
    assert preview_response.status_code == 200
    preview = preview_response.json()["data"]
    generated = client.post("/api/v1/case-docs/generate", json=payload)
    assert generated.status_code == 200
    workbook = load_workbook(BytesIO(generated.content), keep_vba=True)
    try:
        assert preview["cells"]
        assert {cell["target_no"] for cell in preview["cells"]} == {None, 1, 2, 3}
        for cell in preview["cells"]:
            assert "{{" in cell["template_text"]
            assert cell["resolved_text"] == str(workbook[cell["sheet_name"]][cell["cell"]].value or "")
        assert next(cell for cell in preview["cells"] if cell["cell"] == "M6")["resolved_text"] == "10.10.1.10"
        assert next(cell for cell in preview["cells"] if cell["cell"] == "Q6")["resolved_text"] == "10.10.1.11"
        assert next(cell for cell in preview["cells"] if cell["cell"] == "U6")["resolved_text"] == "10.10.1.10"
        unknown = [cell for cell in preview["cells"] if cell["unresolved_placeholders"]]
        assert len(unknown) == 1
        assert unknown[0]["unresolved_placeholders"] == ["{{UNKNOWN_DEVICE}}"]
        assert preview["unresolved_count"] == 1
        assert all(cell["sheet_name"] == workbook.worksheets[0].title for cell in preview["cells"])
    finally:
        workbook.close()
    assert sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == before


def test_unselected_device_stays_unresolved_in_preview(client, monkeypatch, tmp_path):
    context = _export_context(tmp_path, [["GUI_0"]])
    monkeypatch.setattr("app.routers.case_docs.resolve_case_doc_context", lambda *_: context)
    _mock_source_doc_detail(monkeypatch)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["GUI_0"]]))
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["unresolved_count"] == 2
    assert all(cell["target_no"] == 1 for cell in data["cells"] if cell["unresolved_placeholders"])
    assert any(cell["resolved_text"] == "{{SBC_COMMAND_FLOATING_IP}}" for cell in data["cells"])


def test_preview_rejects_ambiguous_values(client, monkeypatch):
    _mock_source_doc_detail(monkeypatch)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["SBC_CL1_0", "SBC_CL1_1"]]))
    assert response.status_code == 400
    assert "SBC_COMMAND_FLOATING_IP" in response.json()["message"]


def test_preview_accepts_slot_qualified_values(client, monkeypatch):
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{SBC_COMMAND_FLOATING_IP__SBC_CL1_0}} / {{SBC_COMMAND_FLOATING_IP__SBC_CL1_1}}"
    source.items[0].rows[1].command_text = "{{TARGET_DEVICE_HOSTNAME}}"
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda *_: source)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["SBC_CL1_0", "SBC_CL1_1"]]))
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["unresolved_count"] == 0
    assert next(cell for cell in data["cells"] if cell["cell"] == "M6")["resolved_text"] == "10.10.1.10 / 10.10.1.11"


@pytest.mark.parametrize("groups", [None, [], [[]], [["missing"]], [["SBC_CL1_0"]] * 21])
def test_preview_requires_explicit_valid_targets(client, groups):
    response = client.post("/api/v1/case-docs/preview", json=_request(client, groups))
    assert response.status_code in (400, 422)


def test_preview_returns_not_found(client, monkeypatch):
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda *_: None)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["SBC_CL1_0"]]))
    assert response.status_code == 404


def test_preview_with_no_body_placeholders_is_empty(client, monkeypatch):
    source = _fake_source_doc_detail(1)
    for row in source.items[0].rows:
        row.command_text = "show status"
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda *_: source)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["SBC_CL1_0"]]))
    assert response.status_code == 200
    assert response.json()["data"]["cells"] == []
    assert response.json()["data"]["unresolved_count"] == 0


def test_preview_database_failure_is_reported(client, monkeypatch):
    def unavailable(*_):
        raise DatabaseConnectionError("Test database unavailable")
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", unavailable)
    response = client.post("/api/v1/case-docs/preview", json=_request(client, [["SBC_CL1_0"]]))
    assert response.status_code == 500


def test_preview_requires_login(client):
    payload = _request(client, [["SBC_CL1_0"]])
    del client.app.dependency_overrides[get_current_user]
    response = client.post("/api/v1/case-docs/preview", json=payload)
    assert response.status_code == 401


def test_preview_is_available_to_non_admin_roles(non_admin_client, monkeypatch):
    _mock_source_doc_detail(monkeypatch)
    response = non_admin_client.post("/api/v1/case-docs/preview", json=_request(non_admin_client, [["SBC_CL1_0"]]))
    assert response.status_code == 200
