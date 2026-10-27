"""Case document generation routes for Sprint 4."""

from typing import Annotated
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response

from app.core.config import AppSettings
from app.core.exceptions import DatabaseConnectionError
from app.core.case_doc_generation import (
    XLSM_MEDIA_TYPE,
    build_case_doc_evidence_workbook_bytes,
    build_case_doc_execution_item_snapshots,
    build_case_doc_workbook_bytes,
    preview_case_doc_workbook,
)
from app.core.responses import (
    ApiResponse,
    CaseDocDataSourceData,
    CaseDocDataSourceListData,
    CaseDocDataSourceRequest,
    CaseDocCommentSearchData,
    CaseDocExecutionBulkUpdateRequest,
    CaseDocExecutionCommentUpdateRequest,
    CaseDocExecutionUpdateRequest,
    CaseDocGenerateRequest,
    CaseDocGenerationPreviewData,
    CaseDocInstanceCompleteRequest,
    CaseDocInstanceCreateRequest,
    CaseDocInstanceDetailData,
    CaseDocInstanceListData,
    CaseDocPreparationUpdateRequest,
    CaseDocMasterOptionsData,
    CaseDocPlaceholderMappingEnabledRequest,
    CaseDocPlaceholderMappingItemData,
    CaseDocPlaceholderMappingListData,
    CaseDocPlaceholderMappingUpsertRequest,
    CaseDocPlaceholderSourceFileListData,
    CaseDocPlaceholderSourcePreviewData,
    CaseDocPlaceholderSampleData,
    CaseDocPlaceholderSampleOptionsData,
    CaseDocPlaceholderSampleRequest,
    CaseDocResolveContextData,
    CaseDocResolveContextRequest,
    CaseDocUnitConfigListData,
    CaseDocUniWorkSelection,
    CaseDocUniCandidatesData,
    RouterEndpointData,
    RouterFoundationData,
    success_response,
)
from app.db.case_docs import (
    create_case_doc_placeholder_mapping,
    get_case_doc_placeholder_source_preview,
    get_case_doc_master_repository,
    list_case_doc_blocks,
    list_case_doc_buildings,
    list_case_doc_placeholder_mappings,
    list_case_doc_placeholder_source_files,
    list_case_doc_prefectures,
    list_case_doc_unit_configs,
    resolve_case_doc_context,
    set_case_doc_placeholder_mapping_enabled,
    update_case_doc_placeholder_mapping,
    validate_case_doc_placeholder_mapping,
)
from app.db.case_doc_instances import (
    complete_case_doc_instance,
    create_case_doc_instance,
    get_case_doc_instance_detail,
    list_case_doc_instances,
    search_case_doc_comments,
    read_case_doc_original_workbook,
    update_case_doc_preparation,
    update_case_doc_execution_comment,
    update_case_doc_execution_item,
    update_case_doc_execution_items,
)
from app.db.source_docs import get_source_doc_detail
from app.db.case_doc_data_sources import list_data_sources, save_data_source, select_data_source
from app.db.case_doc_placeholder_samples import list_placeholder_sample_options, preview_placeholder_sample
from app.routers.health import get_app_settings
from app.core.security import AdminUser, CurrentUser


router = APIRouter(prefix="/case-docs", tags=["case-docs"])


def get_data_source_settings(
    settings: Annotated[AppSettings, Depends(get_app_settings)],
    data_source_id: Annotated[str | None, Query(max_length=40, pattern=r"^[a-z][a-z0-9_-]*$")] = None,
) -> AppSettings:
    try:
        return select_data_source(settings, data_source_id)[0]
    except ValueError as exception:
        raise HTTPException(status_code=400, detail=str(exception)) from exception
    except (OSError, DatabaseConnectionError) as exception:
        raise HTTPException(status_code=503, detail="参照元DBの設定を読み込めませんでした。") from exception


@router.get("/data-sources", response_model=ApiResponse[CaseDocDataSourceListData])
def read_data_sources(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
    include_disabled: bool = False,
) -> ApiResponse[CaseDocDataSourceListData]:
    if include_disabled and current_user.role != "admin":
        raise HTTPException(status_code=403, detail="この操作を実行する権限がありません。")
    try:
        data = list_data_sources(settings, include_disabled)
        if current_user.role != "admin":
            data = data.model_copy(update={"items": [item.model_copy(update={"export_directory": ""}) for item in data.items]})
        return success_response(data, "参照元DB一覧を取得しました。")
    except (OSError, ValueError, DatabaseConnectionError) as exception:
        raise HTTPException(status_code=503, detail="参照元DB一覧を取得できませんでした。") from exception


@router.post("/data-sources", response_model=ApiResponse[CaseDocDataSourceData], status_code=201)
def create_data_source(
    payload: CaseDocDataSourceRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocDataSourceData]:
    del admin_user
    try:
        return success_response(save_data_source(settings, payload, create=True), "参照元DBを登録しました。")
    except ValueError as exception:
        raise HTTPException(status_code=400, detail=str(exception)) from exception
    except (OSError, DatabaseConnectionError) as exception:
        raise HTTPException(status_code=503, detail="参照元DBを登録できませんでした。") from exception


@router.put("/data-sources/{source_id}", response_model=ApiResponse[CaseDocDataSourceData])
def edit_data_source(
    source_id: str,
    payload: CaseDocDataSourceRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocDataSourceData]:
    del admin_user
    if source_id != payload.source_id:
        raise HTTPException(status_code=400, detail="参照元DBの識別子は変更できません。")
    try:
        return success_response(save_data_source(settings, payload, create=False), "参照元DBを更新しました。")
    except ValueError as exception:
        raise HTTPException(status_code=400, detail=str(exception)) from exception
    except (OSError, DatabaseConnectionError) as exception:
        raise HTTPException(status_code=503, detail="参照元DBを更新できませんでした。") from exception


@router.get("/foundation", response_model=ApiResponse[RouterFoundationData])
def read_case_doc_router_foundation() -> ApiResponse[RouterFoundationData]:
    """Return the case document API router foundation status."""

    data = RouterFoundationData(
        resource="case-docs",
        sprint="Sprint 4",
        status="foundation-ready",
        planned_endpoints=[
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/master/blocks",
                purpose="Case document block master options.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/master/prefectures",
                purpose="Case document prefecture master options.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/master/buildings",
                purpose="Case document building master options.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/master/unit-config",
                purpose="Case document unit configuration candidates.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/placeholders",
                purpose="Case document placeholder mappings used during generation.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/placeholders/sources",
                purpose="Selectable source files and columns for placeholder mappings.",
            ),
            RouterEndpointData(
                method="GET",
                path="/api/v1/case-docs/placeholders/source-preview",
                purpose="Filtered source rows for the visual placeholder editor.",
            ),
            RouterEndpointData(
                method="GET", path="/api/v1/case-docs/placeholders/sample-options",
                purpose="Sample units and devices for the placeholder editor.",
            ),
            RouterEndpointData(
                method="POST", path="/api/v1/case-docs/placeholders/sample",
                purpose="Resolve an unsaved definition without writing configuration or case data.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/placeholders",
                purpose="Create a case document placeholder mapping.",
            ),
            RouterEndpointData(
                method="PUT",
                path="/api/v1/case-docs/placeholders/{name}",
                purpose="Update a case document placeholder mapping.",
            ),
            RouterEndpointData(
                method="PATCH",
                path="/api/v1/case-docs/placeholders/{name}/enabled",
                purpose="Enable or disable a case document placeholder mapping.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/placeholders/validate",
                purpose="Validate a placeholder mapping without saving it.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/resolve-context",
                purpose="Resolve case document generation context without manual placeholder input.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/generate",
                purpose="Generate a case document from a source document and resolved context.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/preview",
                purpose="Review actual placeholder substitutions without creating an instance.",
            ),
            RouterEndpointData(
                method="POST",
                path="/api/v1/case-docs/instances",
                purpose="Create a persistent executable case document instance.",
            ),
            RouterEndpointData(
                method="GET", path="/api/v1/case-docs/comments",
                purpose="Search saved comments across case documents without updating them.",
            ),
        ],
    )
    return success_response(data, "Case document API foundation is ready.")


@router.get("/master/prefectures", response_model=ApiResponse[CaseDocMasterOptionsData])
def read_case_doc_prefectures(
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocMasterOptionsData]:
    """Return selectable prefectures for case document generation."""

    try:
        data = list_case_doc_prefectures(settings)
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=400, detail="選択した参照元DBのユニット構成を読み込めませんでした。") from exception
    return success_response(data, "Case document prefectures were retrieved.")


@router.get("/master/blocks", response_model=ApiResponse[CaseDocMasterOptionsData])
def read_case_doc_blocks(
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocMasterOptionsData]:
    """Return selectable blocks for case document generation."""

    try:
        data = list_case_doc_blocks(settings)
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=400, detail="選択した参照元DBのユニット構成を読み込めませんでした。") from exception
    return success_response(data, "ブロック一覧を取得しました。")


@router.get("/master/buildings", response_model=ApiResponse[CaseDocMasterOptionsData])
def read_case_doc_buildings(
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
    prefecture: Annotated[str | None, Query(min_length=1)] = None,
    block: Annotated[str | None, Query(min_length=1)] = None,
) -> ApiResponse[CaseDocMasterOptionsData]:
    """Return selectable buildings for a block or legacy prefecture selection."""

    if prefecture is None and block is None:
        raise HTTPException(status_code=400, detail="ブロックまたは都道府県を指定してください。")
    try:
        data = list_case_doc_buildings(settings, prefecture, block=block)
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=400, detail="選択した参照元DBのユニット構成を読み込めませんでした。") from exception
    return success_response(data, "Case document buildings were retrieved.")


@router.get("/master/unit-config", response_model=ApiResponse[CaseDocUnitConfigListData])
def read_case_doc_unit_configs(
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
    building: Annotated[str, Query(min_length=1)],
    prefecture: Annotated[str | None, Query(min_length=1)] = None,
    block: Annotated[str | None, Query(min_length=1)] = None,
) -> ApiResponse[CaseDocUnitConfigListData]:
    """Return unit configuration candidates for selected location."""

    if prefecture is None and block is None:
        raise HTTPException(status_code=400, detail="ブロックまたは都道府県を指定してください。")
    try:
        data = list_case_doc_unit_configs(settings, prefecture, building, block=block)
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=400, detail="選択した参照元DBのユニット構成を読み込めませんでした。") from exception
    return success_response(data, "Case document unit configurations were retrieved.")


@router.get("/uni/clusters", response_model=ApiResponse[CaseDocMasterOptionsData])
def read_uni_clusters(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
    block: Annotated[str, Query(min_length=1)],
    building: Annotated[str, Query(min_length=1)],
) -> ApiResponse[CaseDocMasterOptionsData]:
    del current_user
    if settings.case_doc_generation_mode != "uni":
        raise HTTPException(status_code=400, detail="選択した参照元DBはUNI方式ではありません。")
    try:
        data = get_case_doc_master_repository(settings).list_clusters(block, building)
    except (OSError, ValueError) as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    return success_response(data, "UNIのクラスタ番号を取得しました。")


@router.post("/uni/target-candidates", response_model=ApiResponse[CaseDocUniCandidatesData])
def read_uni_target_candidates(
    payload: CaseDocUniWorkSelection,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocUniCandidatesData]:
    del current_user
    if settings.case_doc_generation_mode != "uni":
        raise HTTPException(status_code=400, detail="選択した参照元DBはUNI方式ではありません。")
    try:
        data = get_case_doc_master_repository(settings).target_candidates(payload)
    except (OSError, ValueError) as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    return success_response(data, "UNIのホスト名を生成し、抽出ファイルと照合しました。")


@router.get("/placeholders", response_model=ApiResponse[CaseDocPlaceholderMappingListData])
def read_case_doc_placeholder_mappings(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderMappingListData]:
    """Return placeholder mappings used for case document generation."""

    del current_user
    try:
        data = list_case_doc_placeholder_mappings(settings)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception
    return success_response(data, "Case document placeholder mappings were retrieved.")


@router.get("/placeholders/sources", response_model=ApiResponse[CaseDocPlaceholderSourceFileListData])
def read_case_doc_placeholder_source_files(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderSourceFileListData]:
    """Return selectable export files and their column names."""

    del current_user
    try:
        data = list_case_doc_placeholder_source_files(settings)
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    return success_response(data, "Case document placeholder source files were retrieved.")


@router.get("/placeholders/source-preview", response_model=ApiResponse[CaseDocPlaceholderSourcePreviewData])
def read_case_doc_placeholder_source_preview(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
    source_file: Annotated[str, Query(min_length=1)],
    fs_cluster_name: Annotated[str | None, Query()] = None,
    block: Annotated[str | None, Query()] = None,
    prefecture: Annotated[str | None, Query()] = None,
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=200)] = 50,
) -> ApiResponse[CaseDocPlaceholderSourcePreviewData]:
    """Return one export table filtered through unit configuration data."""

    del current_user
    try:
        data = get_case_doc_placeholder_source_preview(
            settings,
            source_file,
            fs_cluster_name,
            block,
            prefecture,
            page,
            page_size,
        )
    except (OSError, ValueError) as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    return success_response(data, "Case document placeholder source preview was retrieved.")


@router.get("/placeholders/sample-options", response_model=ApiResponse[CaseDocPlaceholderSampleOptionsData])
def read_placeholder_sample_options(
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderSampleOptionsData]:
    del admin_user
    try:
        data = list_placeholder_sample_options(settings)
    except (OSError, ValueError, InvalidFileException, BadZipFile) as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    return success_response(data, "Placeholder sample units were retrieved.")


@router.post("/placeholders/sample", response_model=ApiResponse[CaseDocPlaceholderSampleData])
def read_placeholder_sample(
    payload: CaseDocPlaceholderSampleRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderSampleData]:
    del admin_user
    try:
        data = preview_placeholder_sample(settings, payload)
    except (OSError, ValueError, InvalidFileException, BadZipFile) as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    return success_response(data, "Unsaved placeholder sample was evaluated without saving.")


@router.post("/placeholders/validate", response_model=ApiResponse[CaseDocPlaceholderMappingItemData])
def validate_case_doc_placeholder_mapping_payload(
    payload: CaseDocPlaceholderMappingUpsertRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderMappingItemData]:
    """Validate a placeholder mapping without saving it."""

    del admin_user
    try:
        data = validate_case_doc_placeholder_mapping(settings, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception
    return success_response(data, "Case document placeholder mapping is valid.")


@router.post(
    "/placeholders",
    response_model=ApiResponse[CaseDocPlaceholderMappingItemData],
    status_code=status.HTTP_201_CREATED,
)
def create_case_doc_placeholder_mapping_payload(
    payload: CaseDocPlaceholderMappingUpsertRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderMappingItemData]:
    """Create a placeholder mapping used for case document generation."""

    del admin_user
    try:
        data = create_case_doc_placeholder_mapping(settings, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception
    return success_response(data, "Case document placeholder mapping was created.")


@router.put("/placeholders/{name}", response_model=ApiResponse[CaseDocPlaceholderMappingItemData])
def update_case_doc_placeholder_mapping_payload(
    name: str,
    payload: CaseDocPlaceholderMappingUpsertRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderMappingItemData]:
    """Update a placeholder mapping used for case document generation."""

    del admin_user
    try:
        data = update_case_doc_placeholder_mapping(settings, name, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception
    return success_response(data, "Case document placeholder mapping was updated.")


@router.patch("/placeholders/{name}/enabled", response_model=ApiResponse[CaseDocPlaceholderMappingItemData])
def set_case_doc_placeholder_mapping_enabled_payload(
    name: str,
    payload: CaseDocPlaceholderMappingEnabledRequest,
    admin_user: AdminUser,
    settings: Annotated[AppSettings, Depends(get_data_source_settings)],
) -> ApiResponse[CaseDocPlaceholderMappingItemData]:
    """Enable or disable a placeholder mapping used for case document generation."""

    del admin_user
    try:
        data = set_case_doc_placeholder_mapping_enabled(settings, name, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception
    return success_response(data, "Case document placeholder mapping status was updated.")


@router.post("/resolve-context", response_model=ApiResponse[CaseDocResolveContextData])
def resolve_case_doc_generation_context(
    payload: CaseDocResolveContextRequest,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocResolveContextData]:
    """Resolve generation context without accepting manual placeholder values."""

    try:
        data = resolve_case_doc_context(settings, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception

    return success_response(data, "Case document context was resolved.")


@router.post("/preview", response_model=ApiResponse[CaseDocGenerationPreviewData])
def preview_case_doc_generation(
    payload: CaseDocGenerateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocGenerationPreviewData]:
    """Preview the selected targets using the Excel renderer without DB writes."""

    del current_user
    if not payload.target_slot_groups and not payload.uni_targets:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="各targetの対象装置を選択してください。")
    try:
        context = resolve_case_doc_context(settings, payload)
        source_doc = get_source_doc_detail(settings, payload.source_doc_id)
        if source_doc is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="原本が見つかりませんでした。")
        data = preview_case_doc_workbook(context, source_doc)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    except OSError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="置換結果の確認に必要なファイルを読み込めませんでした。") from exception
    return success_response(data, "置換結果を確認しました。")


@router.post("/generate")
def generate_case_doc(
    payload: CaseDocGenerateRequest,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> Response:
    """Generate a downloadable case document workbook."""

    try:
        context = resolve_case_doc_context(settings, payload)
    except ValueError as exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exception),
        ) from exception

    try:
        source_doc = get_source_doc_detail(settings, payload.source_doc_id)
    except DatabaseConnectionError as exception:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=str(exception),
        ) from exception

    if source_doc is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="source document was not found.",
        )

    try:
        workbook_bytes = build_case_doc_workbook_bytes(context, source_doc)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    suffix = context.unit_config.unit_config_id if context.unit_config else context.data_source_id
    filename = f"case-doc-{payload.source_doc_id}-{suffix}.xlsm"
    return Response(
        content=workbook_bytes,
        media_type=XLSM_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post(
    "/instances",
    response_model=ApiResponse[CaseDocInstanceDetailData],
    status_code=status.HTTP_201_CREATED,
)
def create_case_doc_instance_resource(
    payload: CaseDocInstanceCreateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Generate and persist an executable case document instance."""

    try:
        authenticated_payload = payload.model_copy(update={"created_by": current_user.display_name})
        context = resolve_case_doc_context(settings, authenticated_payload)
        source_doc = get_source_doc_detail(settings, payload.source_doc_id)
        if source_doc is None:
            raise ValueError("原本が見つかりませんでした。")
        workbook_bytes = build_case_doc_workbook_bytes(context, source_doc)
        execution_items = build_case_doc_execution_item_snapshots(source_doc, context, workbook_bytes)
        data = create_case_doc_instance(
            settings,
            source_doc,
            context,
            workbook_bytes,
            execution_items,
            current_user.display_name,
            payload.case_name,
        )
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    except (FileNotFoundError, OSError) as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception

    return success_response(data, "案件CSを作成し、実行対象として保存しました。")


@router.get("/instances", response_model=ApiResponse[CaseDocInstanceListData])
def read_case_doc_instances(
    settings: Annotated[AppSettings, Depends(get_app_settings)],
    keyword: Annotated[str | None, Query(min_length=1, max_length=200)] = None,
    tag_paths: Annotated[list[str] | None, Query(alias="tag_path", min_length=1)] = None,
    execution_status: Annotated[
        str | None,
        Query(pattern=r"^(all|active|not_started|in_progress|completed)$"),
    ] = None,
) -> ApiResponse[CaseDocInstanceListData]:
    """Return persistent executable case document instances."""

    try:
        data = list_case_doc_instances(
            settings,
            keyword=keyword,
            tag_paths=tag_paths,
            execution_status=execution_status,
        )
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "案件CS実行一覧を取得しました。")


@router.get("/comments", response_model=ApiResponse[CaseDocCommentSearchData])
def read_case_doc_comments(
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
    keyword: Annotated[str | None, Query(max_length=200)] = None,
    case_keyword: Annotated[str | None, Query(max_length=200)] = None,
    execution_status: Annotated[str | None, Query(pattern=r"^(all|active|not_started|in_progress|completed)$")] = None,
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=100)] = 20,
) -> ApiResponse[CaseDocCommentSearchData]:
    del current_user
    try:
        data = search_case_doc_comments(settings, keyword, case_keyword, execution_status, page, page_size)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "Case document comments were retrieved.")


@router.get("/instances/{case_document_id}", response_model=ApiResponse[CaseDocInstanceDetailData])
def read_case_doc_instance_detail(
    case_document_id: int,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Return one executable case document instance."""

    try:
        data = get_case_doc_instance_detail(settings, case_document_id)
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    if data is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="案件CSが見つかりませんでした。")
    return success_response(data, "案件CS実行詳細を取得しました。")


@router.put(
    "/instances/{case_document_id}/preparation",
    response_model=ApiResponse[CaseDocInstanceDetailData],
)
def update_case_doc_preparation_resource(
    case_document_id: int,
    payload: CaseDocPreparationUpdateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Save preparation values entered before case document execution."""

    try:
        authenticated_payload = payload.model_copy(update={"updated_by": current_user.display_name})
        data = update_case_doc_preparation(settings, case_document_id, authenticated_payload)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "案件CSの工事情報を保存しました。")


@router.patch(
    "/instances/{case_document_id}/items/bulk",
    response_model=ApiResponse[CaseDocInstanceDetailData],
)
def update_case_doc_execution_items_resource(
    case_document_id: int,
    payload: CaseDocExecutionBulkUpdateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Check or skip pending time cells together without overwriting completed cells."""

    try:
        data = update_case_doc_execution_items(
            settings,
            case_document_id,
            payload,
            current_user.display_name,
        )
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "案件CSの実施状態を一括更新しました。")


@router.patch(
    "/instances/{case_document_id}/items/{execution_item_id}",
    response_model=ApiResponse[CaseDocInstanceDetailData],
)
def update_case_doc_execution_item_resource(
    case_document_id: int,
    execution_item_id: int,
    payload: CaseDocExecutionUpdateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Check, skip, or reset one case document time cell."""

    try:
        authenticated_payload = payload.model_copy(update={"performed_by": current_user.display_name})
        data = update_case_doc_execution_item(
            settings,
            case_document_id,
            execution_item_id,
            authenticated_payload,
        )
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "案件CSの実施状態を更新しました。")


@router.put(
    "/instances/{case_document_id}/comments/{group_start_row_order}",
    response_model=ApiResponse[CaseDocInstanceDetailData],
)
def update_case_doc_execution_comment_resource(
    case_document_id: int,
    group_start_row_order: int,
    payload: CaseDocExecutionCommentUpdateRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Save the comment attached to one minor-number execution group."""

    try:
        data = update_case_doc_execution_comment(
            settings,
            case_document_id,
            group_start_row_order,
            payload,
            current_user.display_name,
        )
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "小項番のコメントを保存しました。")


@router.post(
    "/instances/{case_document_id}/complete",
    response_model=ApiResponse[CaseDocInstanceDetailData],
)
def complete_case_doc_instance_resource(
    case_document_id: int,
    payload: CaseDocInstanceCompleteRequest,
    current_user: CurrentUser,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> ApiResponse[CaseDocInstanceDetailData]:
    """Mark a case document instance as completed."""

    del payload, current_user
    try:
        data = complete_case_doc_instance(settings, case_document_id)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception
    return success_response(data, "案件CSを完了しました。")


@router.get("/instances/{case_document_id}/export")
def export_case_doc_instance_evidence(
    case_document_id: int,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> Response:
    """Download a copy of the original xlsm with WebUI execution evidence."""

    try:
        detail = get_case_doc_instance_detail(settings, case_document_id)
        if detail is None:
            raise ValueError("案件CSが見つかりませんでした。")
        original_bytes = read_case_doc_original_workbook(settings, case_document_id)
        evidence_bytes = build_case_doc_evidence_workbook_bytes(original_bytes, detail)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception

    filename = f"{detail.case_document_key}-evidence.xlsm"
    return Response(
        content=evidence_bytes,
        media_type=XLSM_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/instances/{case_document_id}/workbook")
def export_unstarted_case_doc_workbook(
    case_document_id: int,
    settings: Annotated[AppSettings, Depends(get_app_settings)],
) -> Response:
    """Download the original case document before WebUI execution begins."""

    try:
        detail = get_case_doc_instance_detail(settings, case_document_id)
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception

    if detail is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="案件CSが見つかりませんでした。")
    if detail.status != "active" or detail.checked_count + detail.skipped_count > 0:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="実行開始後または完了済みの案件CSは、元の案件CS Excelを出力できません。",
        )

    try:
        workbook_bytes = read_case_doc_original_workbook(settings, case_document_id)
    except ValueError as exception:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exception)) from exception
    except DatabaseConnectionError as exception:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(exception)) from exception

    filename = f"{detail.case_document_key}.xlsm"
    return Response(
        content=workbook_bytes,
        media_type=XLSM_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
