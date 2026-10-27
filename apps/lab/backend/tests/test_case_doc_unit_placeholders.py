"""Unit-scoped definitions use the selected row independently of target devices."""

from io import BytesIO

import pytest
import yaml
from openpyxl import load_workbook

from app.core.case_doc_generation import build_case_doc_workbook_bytes
from app.core.responses import CaseDocResolveContextRequest, ModuleRowDeviceEntryData
from app.db.case_doc_repositories import ExportFileCaseDocMasterRepository, SeedCaseDocMasterRepository
from test_case_doc_placeholder_keys import _write_workbook
from test_case_docs import _fake_source_doc_detail


UNIT_FILE = "\u30e6\u30cb\u30c3\u30c8\u69cb\u6210.xlsx"
SLOTS = ("GUI_0", "GUI_1", "FS_0", "FS_1", "HFS_0", "HFS_1")


def _unit_mapping(slot, **overrides):
    return {
        "name": f"UNIT_{slot}_HOSTNAME", "enabled": True, "scope": "unit",
        "source_file": UNIT_FILE, "value_column": slot + "\u7cfb",
        "source_column": f"unit_{slot.lower()}_hostname", **overrides,
    }


def _write_mappings(path, mappings):
    path.write_text(yaml.safe_dump({"placeholders": mappings}), encoding="utf-8")


@pytest.fixture
def repository(tmp_path):
    # Repeated GUI host names must not cause a value from a different unit to be used.
    _write_workbook(tmp_path / UNIT_FILE, [
        ["unit_config_id", "FS\u30af\u30e9\u30b9\u30bf\u540d", "block", "prefecture", "building", *[slot + "\u7cfb" for slot in SLOTS]],
        ["unit-a", "FS-A", "B1", "Tokyo", "building-a", "GUI-A", "GUI-B", "FS-A-0", "FS-A-1", "HFS-A-0", "HFS-A-1"],
        ["unit-b", "FS-B", "B2", "Tokyo", "building-a", "GUI-A", "GUI-B", "FS-B-0", "FS-B-1", "HFS-B-0", "HFS-B-1"],
    ])
    _write_workbook(tmp_path / "GUI.xlsx", [["host_name", "ip"], ["GUI-A", "192.0.2.10"], ["GUI-B", "192.0.2.11"]])
    mappings = [_unit_mapping(slot) for slot in SLOTS]
    mappings.append({
        "name": "GUI_IP", "enabled": True, "scope": "device", "device_type": "GUI",
        "source_file": "GUI.xlsx", "key_column": "host_name", "value_column": "ip", "source_column": "ip",
    })
    _write_mappings(tmp_path / "mapping.yml", mappings)
    return ExportFileCaseDocMasterRepository(str(tmp_path), str(tmp_path / "mapping.yml"))


def _context(repository, unit_id="unit-a", groups=None):
    return repository.resolve_context(CaseDocResolveContextRequest(
        source_doc_id=1, prefecture="Tokyo", building="building-a", unit_config_id=unit_id,
        target_slot_groups=groups or [["GUI_0", "FS_0"], ["GUI_1", "HFS_1"], ["GUI_1", "HFS_1"]],
    ))


@pytest.mark.parametrize("unit_id, prefix", [("unit-a", "A"), ("unit-b", "B")])
def test_only_selected_unit_values_are_resolved(repository, unit_id, prefix):
    items = {item.placeholder: item for item in _context(repository, unit_id).resolved_placeholders if item.placeholder.startswith("UNIT_")}
    assert len(items) == 6
    assert items["UNIT_FS_1_HOSTNAME"].value == f"FS-{prefix}-1"
    assert items["UNIT_HFS_0_HOSTNAME"].value == f"HFS-{prefix}-0"
    assert all(item.host_name is None for item in items.values())


def test_unit_values_replace_unselected_devices_in_each_target_and_shared_columns(repository):
    source = _fake_source_doc_detail(1)
    row = source.items[0].rows[0]
    row.work_text = "work on {{UNIT_HFS_0_HOSTNAME}}"
    row.expected_result = "verify {{UNIT_GUI_1_HOSTNAME}}"
    row.device_entries = [
        ModuleRowDeviceEntryData(slot_no=1, command_text="{{UNIT_HFS_1_HOSTNAME}}"),
        ModuleRowDeviceEntryData(slot_no=2, command_text="{{UNIT_FS_1_HOSTNAME}}"),
        ModuleRowDeviceEntryData(slot_no=3, command_text="{{UNIT_HFS_1_HOSTNAME}}"),
    ]
    source.items[0].rows[1].command_text = "{{GUI_IP}} / {{TARGET_DEVICE_HOSTNAME}}"
    workbook = load_workbook(BytesIO(build_case_doc_workbook_bytes(_context(repository), source)), keep_vba=True)
    try:
        sheet = workbook.worksheets[0]
        assert [sheet[cell].value for cell in ("M6", "Q6", "U6")] == ["HFS-A-1", "FS-A-1", "HFS-A-1"]
        assert sheet["E6"].value == "work on HFS-A-0"
        assert sheet["I6"].value == "verify GUI-B"
        assert sheet["M7"].value == "192.0.2.10 / GUI-A\nFS-A-0"
        assert sheet["Q7"].value == "192.0.2.11 / GUI-B\nHFS-A-1"
        assert sheet["U7"].value == sheet["Q7"].value
    finally:
        workbook.close()


def test_device_scope_restrictions_remain_while_unit_scope_is_shared(repository):
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{GUI_IP}} / {{UNIT_GUI_1_HOSTNAME}}"
    workbook = load_workbook(BytesIO(build_case_doc_workbook_bytes(_context(repository, groups=[["FS_0"]]), source)), keep_vba=True)
    try:
        assert workbook.worksheets[0]["M6"].value == "{{GUI_IP}} / GUI-B"
    finally:
        workbook.close()


def test_missing_unit_column_is_reported(repository, tmp_path):
    _write_mappings(tmp_path / "mapping.yml", [_unit_mapping("MISSING")])
    with pytest.raises(ValueError, match="UNIT_MISSING_HOSTNAME"):
        _context(repository)


def test_blank_value_does_not_fall_back_to_another_unit(repository, tmp_path):
    _write_workbook(tmp_path / UNIT_FILE, [
        ["unit_config_id", "fs_cluster_name", "block", "prefecture", "building", "GUI_0\u7cfb", "HFS_0\u7cfb"],
        ["unit-a", "FS-A", "B1", "Tokyo", "building-a", "GUI-A", ""],
        ["unit-b", "FS-B", "B2", "Tokyo", "building-a", "GUI-A", "other-unit-value"],
    ])
    _write_mappings(tmp_path / "mapping.yml", [_unit_mapping("HFS_0")])
    assert _context(repository, groups=[["GUI_0"]]).resolved_placeholders == []


def test_disabled_unit_mapping_is_not_resolved(repository, tmp_path):
    _write_mappings(tmp_path / "mapping.yml", [_unit_mapping("MISSING", enabled=False)])
    assert _context(repository).resolved_placeholders == []


def test_scope_is_explicit_not_inferred_from_placeholder_name(repository, tmp_path):
    _write_mappings(tmp_path / "mapping.yml", [_unit_mapping("HFS_1", name="FIXED_HOST")])
    items = _context(repository, groups=[["FS_0"]]).resolved_placeholders
    assert [(item.placeholder, item.value, item.host_name) for item in items] == [("FIXED_HOST", "HFS-A-1", None)]


def test_unit_scope_works_in_seed_mode(tmp_path):
    path = tmp_path / "mapping.yml"
    _write_mappings(path, [_unit_mapping("GUI_1")])
    repository = SeedCaseDocMasterRepository(str(path))
    unit = repository.list_unit_configs(repository.list_prefectures().items[0].value, repository.list_buildings(repository.list_prefectures().items[0].value).items[0].value).items[0]
    context = repository.resolve_context(CaseDocResolveContextRequest(
        source_doc_id=1, prefecture=unit.prefecture, building=unit.building,
        unit_config_id=unit.unit_config_id, target_slot_groups=[["SBC_CL1_0"]],
    ))
    assert context.resolved_placeholders[0].placeholder == "UNIT_GUI_1_HOSTNAME"
    assert context.resolved_placeholders[0].host_name is None


def _unit_payload(**overrides):
    return {"name": "UNIT_TEST", "enabled": True, "scope": "unit", "source_file": UNIT_FILE,
            "value_column": "GUI_0\u7cfb", "source_column": "unit_test", **overrides}


def test_api_accepts_unit_mapping_without_key_or_device(client, test_settings, tmp_path):
    path = tmp_path / "mapping.yml"
    _write_mappings(path, [])
    test_settings.case_doc_placeholder_mapping_path = str(path)
    response = client.post("/api/v1/case-docs/placeholders", json=_unit_payload())
    assert response.status_code == 201
    assert response.json()["data"]["scope"] == "unit"
    assert response.json()["data"]["key_column"] == ""
    listed = client.get("/api/v1/case-docs/placeholders").json()["data"]["items"]
    assert listed[0]["scope"] == "unit"
    edited = client.put("/api/v1/case-docs/placeholders/UNIT_TEST", json=_unit_payload(value_column="GUI_1\u7cfb"))
    assert edited.status_code == 200
    assert edited.json()["data"]["value_column"] == "GUI_1\u7cfb"
    toggled = client.patch("/api/v1/case-docs/placeholders/UNIT_TEST/enabled", json={"enabled": False})
    assert toggled.status_code == 200


@pytest.mark.parametrize("override", [
    {"source_file": "GUI.xlsx"}, {"source_file": "../unit_config.xlsx"},
    {"source_device_type": "GUI"}, {"key_value": "constant"},
    {"scope": "device", "device_type": "GUI"}, {"scope": "common", "key_value": "constant"},
])
def test_api_rejects_invalid_scope_settings_without_changing_config(client, test_settings, tmp_path, override):
    path = tmp_path / "mapping.yml"
    _write_mappings(path, [])
    test_settings.case_doc_placeholder_mapping_path = str(path)
    before = path.read_bytes()
    response = client.post("/api/v1/case-docs/placeholders", json=_unit_payload(**override))
    assert response.status_code == 400
    assert path.read_bytes() == before


@pytest.mark.parametrize("override", [
    {"source_file": "GUI.xlsx"}, {"source_device_type": "GUI"}, {"key_value": "constant"},
])
def test_unit_validation_errors_use_current_display_name(client, test_settings, tmp_path, override):
    path = tmp_path / "mapping.yml"
    _write_mappings(path, [])
    test_settings.case_doc_placeholder_mapping_path = str(path)
    before = path.read_bytes()
    response = client.post("/api/v1/case-docs/placeholders", json=_unit_payload(**override))
    assert response.status_code == 400
    assert "\u30e6\u30cb\u30c3\u30c8\u69cb\u6210\u306e\u5024" in response.json()["message"]
    assert path.read_bytes() == before


def test_preview_keeps_unit_mapping_on_value_column(client, test_settings, tmp_path, repository):
    directory = tmp_path / "ssc"
    directory.mkdir()
    (directory / UNIT_FILE).write_bytes((tmp_path / UNIT_FILE).read_bytes())
    test_settings.case_doc_master_source = "export_file"
    test_settings.case_doc_access_export_dir = str(tmp_path)
    test_settings.case_doc_placeholder_mapping_path = str(tmp_path / "mapping.yml")
    response = client.get("/api/v1/case-docs/placeholders/source-preview", params={"source_file": UNIT_FILE})
    assert response.status_code == 200
    data = response.json()["data"]
    assert len([item for item in data["mappings"] if item["scope"] == "unit"]) == 6
    assert data["total_count"] == 2
