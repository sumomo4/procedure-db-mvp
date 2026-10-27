"""Grouped targets retain independent CS columns and shared per-column checks."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.case_doc_generation import build_case_doc_execution_item_snapshots, build_case_doc_workbook_bytes
from app.core.responses import CaseDocResolveContextRequest
from app.db.case_doc_repositories import ExportFileCaseDocMasterRepository
from test_case_docs import _fake_source_doc_detail, _mock_source_doc_detail, _tokyo_building, _tokyo_prefecture


def _request(client, groups):
    return {
        "source_doc_id": 1,
        "prefecture": _tokyo_prefecture(client),
        "building": _tokyo_building(client),
        "unit_config_id": "unit-tokyo-001",
        "target_slot_groups": groups,
    }


def test_same_device_can_be_selected_in_separate_targets(client, monkeypatch):
    _mock_source_doc_detail(monkeypatch)
    payload = _request(client, [["SBC_CL1_0"], ["SBC_CL1_0"], ["SBC_CL1_1"]])
    resolved = client.post("/api/v1/case-docs/resolve-context", json=payload)
    assert resolved.status_code == 200
    assert [group[0]["slot_key"] for group in resolved.json()["data"]["target_groups"]] == [
        "SBC_CL1_0", "SBC_CL1_0", "SBC_CL1_1"
    ]
    generated = client.post("/api/v1/case-docs/generate", json=payload)
    assert generated.status_code == 200
    workbook = load_workbook(BytesIO(generated.content), keep_vba=True)
    sheet = workbook.worksheets[0]
    assert [sheet.cell(4, column).value for column in (13, 17, 21)] == [
        "sbc-tyo-cl1-0", "sbc-tyo-cl1-0", "sbc-tyo-cl1-1"
    ]
    assert [sheet.cell(6, column).value for column in (13, 17, 21)] == ["10.10.1.10", "10.10.1.10", "10.10.1.11"]


@pytest.mark.parametrize("groups", [[], [[]], [["missing"]], [["SBC_CL1_0"]] * 21])
def test_invalid_target_groups_are_rejected(client, groups):
    response = client.post("/api/v1/case-docs/resolve-context", json=_request(client, groups))
    assert response.status_code in (400, 422)


def test_duplicates_within_one_target_are_collapsed(client):
    response = client.post("/api/v1/case-docs/resolve-context", json=_request(client, [["SBC_CL1_0", "SBC_CL1_0"]]))
    assert response.status_code == 200
    assert len(response.json()["data"]["target_groups"][0]) == 1


def _write_export(path, columns, rows):
    workbook = Workbook()
    workbook.active.append(columns)
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


def _export_context(tmp_path, groups=None):
    _write_export(tmp_path / "unit_config.xlsx", [
        "fs_cluster_name", "block", "prefecture", "building", "FS_0", "GUI_0", "HOST", "SBC_CL1_0"
    ], [["FS-cluster", "block-1", "Tokyo", "test-building", "FS-B", "GUI-A", "HOST-A", "SBC-A"]])
    _write_export(tmp_path / "FS.xlsx", ["host_name"], [["FS-B"]])
    _write_export(tmp_path / "GUI.xlsx", ["host_name", "EMS\u30b3\u30de\u30f3\u30c9\u7528IP\u30a2\u30c9\u30ec\u30b9"], [["GUI-A", "192.0.2.10"]])
    _write_export(tmp_path / "SBC.xlsx", ["host_name", "command_floating_ip"], [["SBC-A", "192.0.2.20"]])
    _write_export(tmp_path / "HSS.xlsx", ["host_name"], [["HSS-A"]])
    repository = ExportFileCaseDocMasterRepository(str(tmp_path))
    payload = CaseDocResolveContextRequest(
        source_doc_id=1, prefecture="Tokyo", building="test-building", block="block-1", target_slot_groups=groups
    )
    return repository.resolve_context(payload)


def test_all_nonexcluded_unit_columns_are_candidates(tmp_path):
    context = _export_context(tmp_path)
    assert {item.column_name for item in context.host_assignments} == {"fs_cluster_name", "FS_0", "GUI_0", "HOST", "SBC_CL1_0"}
    assert all(item.column_name not in {"block", "prefecture", "building"} for item in context.host_assignments)


def test_mixed_devices_share_one_target_without_adding_check_cells(tmp_path):
    context = _export_context(tmp_path, [["FS_0", "GUI_0", "HOST"], ["GUI_0"]])
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{GUI_EMS_COMMAND_IP}}"
    source.items[0].rows[1].command_text = "{{TARGET_DEVICE_HOSTNAME}}"
    content = build_case_doc_workbook_bytes(context, source)
    sheet = load_workbook(BytesIO(content), keep_vba=True).worksheets[0]
    assert [sheet.cell(row, 13).value for row in (2, 3, 4)] == ["FS-B", "GUI-A", "HOST-A"]
    assert sheet["M6"].value == "192.0.2.10"
    assert sheet["M7"].value == "FS-B\nGUI-A\nHOST-A"
    assert context.target_device_slots[0].host_name == "FS-B\nGUI-A\nHOST-A"
    snapshots = build_case_doc_execution_item_snapshots(source, context, content)
    assert len(snapshots) == 4
    assert [item["target_no"] for item in snapshots] == [1, 2, 1, 2]
    assert snapshots[0]["host_name"] == "FS-B\nGUI-A\nHOST-A"
    assert snapshots[0]["excel_cell"] == "J6"
    assert snapshots[1]["excel_cell"] == "N6"


def test_ambiguous_placeholder_does_not_silently_choose_a_device(client, monkeypatch):
    _mock_source_doc_detail(monkeypatch)
    response = client.post("/api/v1/case-docs/generate", json=_request(client, [["SBC_CL1_0", "SBC_CL1_1"]]))
    assert response.status_code == 400
    assert "SBC_COMMAND_FLOATING_IP" in response.json()["message"]


def test_slot_qualified_placeholders_can_resolve_distinct_values(client, monkeypatch):
    source = _fake_source_doc_detail(1)
    source.items[0].rows[0].command_text = "{{SBC_COMMAND_FLOATING_IP__SBC_CL1_0}} / {{SBC_COMMAND_FLOATING_IP__SBC_CL1_1}}"
    source.items[0].rows[1].command_text = "unchanged"
    monkeypatch.setattr("app.routers.case_docs.get_source_doc_detail", lambda settings, source_doc_id: source)
    response = client.post("/api/v1/case-docs/generate", json=_request(client, [["SBC_CL1_0", "SBC_CL1_1"]]))
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), keep_vba=True).worksheets[0]
    assert sheet["M6"].value == "10.10.1.10 / 10.10.1.11"


def test_old_unqualified_placeholder_does_not_use_an_unselected_device(tmp_path):
    context = _export_context(tmp_path, [["GUI_0"]])
    source = _fake_source_doc_detail(1)
    content = build_case_doc_workbook_bytes(context, source)
    sheet = load_workbook(BytesIO(content), keep_vba=True).worksheets[0]
    assert sheet["M6"].value == "{{SBC_COMMAND_FLOATING_IP}}"
