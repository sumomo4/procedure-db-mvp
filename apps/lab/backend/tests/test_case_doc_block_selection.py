"""Block-first case selection without changing the stored prefecture or lookup rules."""

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from app.core.responses import CaseDocDataSourceListData, CaseDocResolveContextRequest
from app.db import case_doc_data_sources as sources
from app.db.case_doc_repositories import ExportFileCaseDocMasterRepository, SeedCaseDocMasterRepository
from app.db.case_docs import resolve_case_doc_context
from app.routers.health import get_app_settings


def write_book(path: Path, rows: list[list[str]]) -> None:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


@pytest.fixture
def block_settings(tmp_path, test_settings, monkeypatch):
    root = tmp_path / "exports"
    mapping = tmp_path / "mapping.yml"
    mapping.write_text(yaml.safe_dump({"version": 1, "placeholders": [{
        "name": "GUI_ADDRESS", "enabled": True, "scope": "device", "device_type": "GUI",
        "source_file": "GUI.xlsx", "key_column": "host_name", "value_column": "ip", "source_column": "ip",
    }]}), encoding="utf-8")
    for source_id, host_prefix in [("ssc", "ssc"), ("uni", "uni")]:
        directory = root / source_id
        directory.mkdir(parents=True)
        write_book(directory / "unit_config.xlsx", [
            ["unit_config_id", "fs_cluster_name", "block", "prefecture", "building", "GUI_0"],
            ["unit-osaka", "Same cluster", "West", "Osaka", "Shared building", f"{host_prefix}-osaka"],
            ["unit-kyoto", "Same cluster", "West", "Kyoto", "Shared building", f"{host_prefix}-kyoto"],
            ["unit-tokyo", "Same cluster", "East", "Tokyo", "Shared building", f"{host_prefix}-tokyo"],
            ["unit-east", "East cluster", "East", "Tokyo", "East building", f"{host_prefix}-east"],
        ])
        write_book(directory / "GUI.xlsx", [
            ["host_name", "ip"], [f"{host_prefix}-osaka", "192.0.2.10"],
            [f"{host_prefix}-kyoto", "192.0.2.11"], [f"{host_prefix}-tokyo", "192.0.2.12"],
        ])
        if source_id == "uni":
            target = tmp_path / "placeholder_mappings" / "uni.yml"
            target.parent.mkdir()
            target.write_bytes(mapping.read_bytes())
    settings = test_settings.model_copy(update={"case_doc_master_source": "export_file",
        "case_doc_access_export_dir": str(root), "case_doc_placeholder_mapping_path": str(mapping)})
    items = [sources._to_data(settings, {"source_id": key, "name": key.upper(), "enabled": True, "description": ""}) for key in ("ssc", "uni")]
    monkeypatch.setattr(sources, "list_data_sources", lambda *_args, **_kwargs: CaseDocDataSourceListData(items=items))
    return settings


def test_block_repository_keeps_all_prefectures_and_filters_units(block_settings):
    repository = ExportFileCaseDocMasterRepository(Path(block_settings.case_doc_access_export_dir) / "ssc", block_settings.case_doc_placeholder_mapping_path)
    assert [item.value for item in repository.list_blocks().items] == ["East", "West"]
    assert [item.value for item in repository.list_buildings(block="West").items] == ["Shared building"]
    assert [item.value for item in repository.list_buildings(block="East").items] == ["East building", "Shared building"]
    units = repository.list_unit_configs(None, "Shared building", block="West").items
    assert {item.unit_config_id for item in units} == {"unit-osaka", "unit-kyoto"}
    assert {item.prefecture for item in units} == {"Osaka", "Kyoto"}
    assert repository.list_buildings(block="Unknown").items == []
    assert repository.list_unit_configs(None, "Shared building", block="Unknown").items == []


def test_legacy_prefecture_and_combined_filters_still_work(block_settings):
    repository = ExportFileCaseDocMasterRepository(Path(block_settings.case_doc_access_export_dir) / "ssc", block_settings.case_doc_placeholder_mapping_path)
    assert [item.value for item in repository.list_buildings("Osaka").items] == ["Shared building"]
    assert [item.unit_config_id for item in repository.list_unit_configs("Osaka", "Shared building").items] == ["unit-osaka"]
    assert repository.list_buildings("Tokyo", block="West").items == []
    assert repository.list_unit_configs("Tokyo", "Shared building", block="West").items == []


def test_seed_repository_supports_block_selection():
    repository = SeedCaseDocMasterRepository()
    for block in repository.list_blocks().items:
        buildings = repository.list_buildings(block=block.value).items
        assert buildings
        for building in buildings:
            assert all(unit.block == block.value for unit in repository.list_unit_configs(None, building.value, block=block.value).items)


@pytest.mark.parametrize("source_id", ["ssc", "uni"])
def test_block_routes_are_source_scoped(client, block_settings, source_id):
    client.app.dependency_overrides[get_app_settings] = lambda: block_settings
    params = {"data_source_id": source_id}
    response = client.get("/api/v1/case-docs/master/blocks", params=params)
    assert response.status_code == 200
    assert [item["value"] for item in response.json()["data"]["items"]] == ["East", "West"]
    params["block"] = "West"
    response = client.get("/api/v1/case-docs/master/buildings", params=params)
    assert response.status_code == 200
    assert [item["value"] for item in response.json()["data"]["items"]] == ["Shared building"]
    params["building"] = "Shared building"
    response = client.get("/api/v1/case-docs/master/unit-config", params=params)
    assert response.status_code == 200
    assert {item["prefecture"] for item in response.json()["data"]["items"]} == {"Osaka", "Kyoto"}
    params["prefecture"] = "Kyoto"
    assert [item["unit_config_id"] for item in client.get("/api/v1/case-docs/master/unit-config", params=params).json()["data"]["items"]] == ["unit-kyoto"]


@pytest.mark.parametrize("endpoint,params", [
    ("buildings", {}), ("unit-config", {"building": "Shared building"}),
    ("buildings", {"block": ""}), ("unit-config", {"block": "West", "building": ""}),
])
def test_master_routes_require_a_nonempty_location_filter(client, endpoint, params):
    assert client.get(f"/api/v1/case-docs/master/{endpoint}", params=params).status_code == 400


@pytest.mark.parametrize("endpoint", ["blocks", "buildings?block=West", "unit-config?block=West&building=Shared%20building"])
def test_missing_export_has_a_clear_error(client, test_settings, tmp_path, endpoint):
    settings = test_settings.model_copy(update={"case_doc_master_source": "export_file", "case_doc_access_export_dir": str(tmp_path)})
    client.app.dependency_overrides[get_app_settings] = lambda: settings
    response = client.get(f"/api/v1/case-docs/master/{endpoint}")
    assert response.status_code == 400
    assert "ユニット構成" in response.json()["message"]


@pytest.mark.parametrize("unit_id,prefecture,address", [("unit-osaka", "Osaka", "192.0.2.10"), ("unit-kyoto", "Kyoto", "192.0.2.11")])
def test_selected_unit_preserves_prefecture_and_resolves_the_right_host(block_settings, unit_id, prefecture, address):
    request = CaseDocResolveContextRequest(data_source_id="ssc", source_doc_id=1, block="West", building="Shared building",
        prefecture=prefecture, unit_config_id=unit_id, target_slot_groups=[["GUI_0"], ["GUI_0"]])
    context = resolve_case_doc_context(block_settings, request)
    assert context.unit_config.prefecture == prefecture
    assert context.unit_config.block == "West"
    assert context.target_groups[0] == context.target_groups[1]
    assert next(item.value for item in context.resolved_placeholders if item.placeholder == "GUI_ADDRESS") == address


def test_block_cannot_be_used_as_prefecture_or_cross_into_another_block(block_settings):
    for prefecture, block in [("West", "West"), ("Kyoto", "East")]:
        with pytest.raises(ValueError, match="not found"):
            resolve_case_doc_context(block_settings, CaseDocResolveContextRequest(data_source_id="ssc", source_doc_id=1,
                block=block, prefecture=prefecture, building="Shared building", unit_config_id="unit-kyoto"))
