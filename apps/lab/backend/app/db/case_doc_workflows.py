"""Explicit per-source workflow configuration and UNI hostname calculation."""

from dataclasses import dataclass
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import yaml

from app.core.responses import CaseDocHostAssignmentData

UNI_EXPORT_FILES = ("TTS.xlsx", "収容エリア.xlsx", "装置追加.xlsx", "TTS_定義.xlsx")


def workflow_path(mapping_path: str, source_id: str) -> Path:
    if not re.fullmatch(r"[a-z][a-z0-9_-]{0,39}", source_id):
        raise ValueError("参照元DBの識別子が不正です。")
    root = Path(mapping_path).resolve().parent / "case_doc_workflows"
    path = root / f"{source_id}.yml"
    if path.resolve().parent != root.resolve():
        raise ValueError("案件化方式の設定先が不正です。")
    return path


def read_workflow(path: Path) -> dict:
    if not path.is_file():
        return {"version": 1, "mode": "unit_config"}
    if path.stat().st_size > 65536:
        raise ValueError("案件化方式の設定ファイルが大きすぎます。")
    data = yaml.safe_load(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict) or data.get("version") != 1 or data.get("mode") not in {"unit_config", "uni"}:
        raise ValueError("案件化方式の設定が不正です。")
    if data["mode"] == "uni":
        definition_path(path, data)
    return data


def definition_path(profile_path: Path, profile: dict) -> Path:
    name = profile.get("definition_file")
    if not isinstance(name, str) or not name or Path(name).name != name or any(char in name for char in "/\\:"):
        raise ValueError("UNIの命名規則には同じ設定フォルダ内のExcelを指定してください。")
    path = profile_path.parent / name
    if path.suffix.casefold() != ".xlsx" or path.resolve().parent != profile_path.parent.resolve():
        raise ValueError("UNIの命名規則ファイルの指定が不正です。")
    return path


@dataclass(frozen=True)
class UniBuilding:
    block: str
    prefecture: str
    building: str
    prefix: str
    pair_building: str
    pair_prefix: str


@dataclass
class UniNamingMaster:
    buildings: list[UniBuilding]
    areas: dict[str, dict[int, str]]


def load_uni_master(profile_path: Path) -> UniNamingMaster:
    profile = read_workflow(profile_path)
    if profile["mode"] != "uni":
        raise ValueError("選択した参照元DBはUNI方式ではありません。")
    path = definition_path(profile_path, profile)
    if not path.is_file():
        raise ValueError("UNIのビル別ホスト名規則が未配置です。")
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("UNIの命名規則ファイルが大きすぎます。")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
        raise ValueError("UNIの命名規則Excelを読み込めませんでした。") from error
    try:
        if "定義" not in workbook.sheetnames:
            raise ValueError("UNIの命名規則に「定義」シートがありません。")
        sheet = workbook["定義"]
        if [sheet.cell(3, col).value for col in (2, 3, 5, 6)] != ["ブロック", "府県", "設置ビル", "ホスト名規則"]:
            raise ValueError("UNIの命名規則の見出しが一致しません。")
        buildings, seen = [], set()
        for row in sheet.iter_rows(min_row=4, max_row=24, min_col=2, max_col=8, values_only=True):
            block, prefecture, _, building, prefix, pair, pair_prefix = [str(value).strip() if value is not None else "" for value in row]
            if not any((block, prefecture, building, prefix)):
                continue
            if not all((block, prefecture, building, prefix)) or not re.fullmatch(r"[a-z0-9-]{1,64}", prefix):
                raise ValueError("UNIのビル別ホスト名規則に未設定または不正な値があります。")
            if (block, building) in seen:
                raise ValueError("UNIの命名規則に同じブロック・ビルが重複しています。")
            seen.add((block, building))
            buildings.append(UniBuilding(block, prefecture, building, prefix, pair, pair_prefix))
        if not buildings:
            raise ValueError("UNIのビル別ホスト名規則が未設定です。")
        if sheet.cell(3, 19).value != "収容エリア":
            raise ValueError("UNIの収容エリアの見出しが一致しません。")
        columns = {}
        for col in range(20, min(sheet.max_column, 55) + 1):
            number = sheet.cell(3, col).value
            if isinstance(number, (int, float)) and int(number) == number and 1 <= number <= 36:
                if int(number) in columns:
                    raise ValueError("UNIのクラスタ番号が重複しています。")
                columns[int(number)] = col
        areas = {}
        for row in range(4, 10):
            block = str(sheet.cell(row, 19).value or "").strip()
            if not block:
                continue
            if block in areas:
                raise ValueError("UNIの収容エリアに同じブロックが重複しています。")
            areas[block] = {}
            for number, col in columns.items():
                area = str(sheet.cell(row, col).value or "").strip()
                if area and area not in {"-", "‐", "ー", "―", "−", "?"}:
                    areas[block][number] = area
        return UniNamingMaster(buildings, areas)
    finally:
        workbook.close()


def generate_uni_hosts(building: UniBuilding, cluster_no: int, *, target_no: int = 1) -> list[CaseDocHostAssignmentData]:
    if type(cluster_no) is not int or type(target_no) is not int or not 1 <= cluster_no <= 36 or not 1 <= target_no <= 20:
        raise ValueError("UNIのクラスタ番号またはtarget番号が不正です。")
    first_rule = building.block == "関西1"
    group = (cluster_no - 1) // 4 + 1 if first_rule or cluster_no <= 16 else 5 + (cluster_no - 17) // 6
    position = (cluster_no - 1) % 4 + 1 if first_rule or cluster_no <= 16 else (cluster_no - 17) % 6 + 1
    pairs = [
        ("USBE", "0", f"usbe{group}00000"), ("USBE", "1", f"usbe{group}00100"),
        ("USBH", "0", f"usbh{group}{position}0000"), ("USBH", "1", f"usbh{group}{position}0100"),
        ("USBM", "0", f"usbm{group}{position}1000"), ("USBM", "1", f"usbm{group}{position}1100"),
    ]
    if first_rule:
        suffix = f"usce{'00' if cluster_no > 18 else '000'}{(cluster_no - 1) // 2 + 1}00"
        kind = "USCE"
    elif cluster_no <= 16:
        suffix, kind = f"usce000{(cluster_no - 1) // 2 + 1}00", "USCE"
    elif cluster_no < 19:
        suffix, kind = f"usle000{(cluster_no - 1) // 2 + 1}00", "USLE"
    elif cluster_no == 19:
        suffix, kind = f"usle000{(cluster_no - 1) // 2}00", "USLE"
    else:
        suffix, kind = f"usle00{10 + (cluster_no - 20) // 3}00", "USLE"
    pairs.append((kind, None, suffix))
    return [CaseDocHostAssignmentData(slot_key=f"{kind}_T{target_no}" + (f"_{system}" if system is not None else ""),
        device_type=kind, system=system, host_name=building.prefix + suffix,
        column_name=kind + (f"_{system}系" if system is not None else "")) for kind, system, suffix in pairs]
